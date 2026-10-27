"""Converte docs/relatorio-consolidado.md em DOCX (Tahoma, #244061), sem timbrado.

Rodar: python entregas/gerar_docx.py
# ponytail: conversor só do Markdown que o relatório usa (títulos, listas, tabelas, **negrito**, `código`, links)
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

RAIZ = Path(__file__).resolve().parents[1]
ORIGEM = RAIZ / "docs" / "relatorio-consolidado.md"
DESTINO = Path(__file__).with_name("relatorio-consolidado-2026-10-27.docx")
AZUL = RGBColor(0x24, 0x40, 0x61)
FONTE = "Tahoma"


def fonte(estilo, tamanho, negrito=False):
    estilo.font.name = FONTE
    estilo.font.size = Pt(tamanho)
    estilo.font.bold = negrito
    estilo.font.color.rgb = AZUL
    rf = estilo.element.rPr.rFonts
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):  # títulos herdam a fonte do tema
        rf.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), FONTE)


def inline(par, texto):
    # links viram só o texto; **negrito** e `código` viram runs
    texto = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", texto)
    for parte in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", texto):
        if not parte:
            continue
        if parte.startswith("**"):
            r = par.add_run(parte[2:-2]); r.bold = True
        elif parte.startswith("`"):
            r = par.add_run(parte[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9)
        else:
            r = par.add_run(parte)
        if r.font.color.rgb is None:
            r.font.color.rgb = AZUL


def sombrear(celula, cor):
    tc = celula._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), cor)
    tc.append(sh)


def tabela(doc, linhas):
    cel = [[c.strip() for c in l.strip().strip("|").split("|")] for l in linhas if not re.match(r"^\|[\s\-|:]+\|$", l.strip())]
    if cel[0][-1] == "Diário":  # no site é link para o diário; no DOCX não leva a lugar nenhum
        cel = [linha[:-1] for linha in cel]
    t = doc.add_table(rows=len(cel), cols=len(cel[0]))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    # largura pelo maior texto de cada coluna (raiz para não esmagar colunas curtas), total 16 cm
    pesos = [max(len(linha[j]) for linha in cel) ** 0.5 for j in range(len(cel[0]))]
    larguras = [Cm(16 * p / sum(pesos)) for p in pesos]
    for i, linha in enumerate(cel):
        for j, txt in enumerate(linha):
            c = t.cell(i, j); c.text = ""; c.width = larguras[j]
            p = c.paragraphs[0]
            inline(p, f"**{txt}**" if i == 0 and txt else txt)
            for r in p.runs:
                r.font.size = Pt(8.5)
            if i == 0:
                sombrear(c, "DCE6F1")
    doc.add_paragraph()


def gerar():
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.5); s.top_margin = s.bottom_margin = Cm(2)
    fonte(doc.styles["Normal"], 10.5)
    for n, tam in ((1, 16), (2, 13), (3, 11.5)):
        fonte(doc.styles[f"Heading {n}"], tam, negrito=True)
    for nome in ("List Bullet", "List Number"):
        fonte(doc.styles[nome], 10.5)

    linhas = ORIGEM.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(linhas):
        l = linhas[i]
        if l.startswith("|"):
            bloco = []
            while i < len(linhas) and linhas[i].startswith("|"):
                bloco.append(linhas[i]); i += 1
            tabela(doc, bloco); continue
        m = re.match(r"^(#{1,3}) (.*)", l)
        if m:
            inline(doc.add_heading("", level=len(m.group(1))), m.group(2))
        elif re.match(r"^\s*- ", l):
            texto = l.strip()[2:]
            while i + 1 < len(linhas) and linhas[i + 1].startswith("  ") and not linhas[i + 1].strip().startswith("- "):
                i += 1; texto += " " + linhas[i].strip()
            inline(doc.add_paragraph(style="List Bullet"), texto)
        elif re.match(r"^\d+\. ", l):
            texto = re.sub(r"^\d+\. ", "", l)
            while i + 1 < len(linhas) and linhas[i + 1].startswith("   "):
                i += 1; texto += " " + linhas[i].strip()
            inline(doc.add_paragraph(style="List Number"), texto)
        elif l.strip():
            texto = l.strip()
            while i + 1 < len(linhas) and linhas[i + 1].strip() and not re.match(r"^(#|\||\s*- |\d+\. |\*\*)", linhas[i + 1]):
                i += 1; texto += " " + linhas[i].strip()
            inline(doc.add_paragraph(), texto)
        i += 1
    doc.save(DESTINO)
    return DESTINO


if __name__ == "__main__":
    print(gerar())
